import datetime
import io

import openpyxl

from app.parsers.attendance import parse_weekly_attendance
from app.parsers.pcc_batch import parse_pcc_batch
from app.parsers.rate_master import parse_rate_master

MON = datetime.date(2026, 9, 14)


def _xlsx_bytes(rows, headers=("Name", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_weekly_attendance_basic():
    data = _xlsx_bytes([
        ("Adams", "X", "", "X", "", "X"),
        ("Baker", "", "", "", "", ""),
    ])
    result = parse_weekly_attendance(data, MON)
    assert not result.warnings
    assert len(result.rows) == 2
    adams = next(r for r in result.rows if r.participant_name == "Adams")
    assert adams.attendance_by_date[MON] is True
    assert adams.attendance_by_date[MON + datetime.timedelta(days=1)] is False
    assert adams.attendance_by_date[MON + datetime.timedelta(days=2)] is True


def test_parse_weekly_attendance_stops_at_trials_section():
    """Trial participants are never billed in PCC (per the source
    process), so anything listed under a "Trials:" marker must be
    excluded -- otherwise they'd permanently show as a false
    MISSING_FROM_BATCH exception."""
    data = _xlsx_bytes([
        ("Adams", "X", "", "", "", ""),
        ("Trials:", "", "", "", "", ""),
        ("NewPerson", "X", "", "", "", ""),
    ])
    result = parse_weekly_attendance(data, MON)
    names = [r.participant_name for r in result.rows]
    assert names == ["Adams"]


def test_parse_weekly_attendance_external_id_preserved_in_name():
    data = _xlsx_bytes([
        ("Arocho, Leonidas (PAM-6)", "X", "", "", "", ""),
    ])
    result = parse_weekly_attendance(data, MON)
    assert result.rows[0].participant_name == "Arocho, Leonidas (PAM-6)"


def test_parse_weekly_attendance_missing_header_warns():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Just some", "random", "columns"])
    ws.append(["a", "b", "c"])
    buf = io.BytesIO()
    wb.save(buf)
    result = parse_weekly_attendance(buf.getvalue(), MON)
    assert result.rows == []
    assert result.warnings


def test_parse_rate_master_basic():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Participant Name", "Payer Source", "Rate", "Exception Notes"])
    ws.append(["Baker", "Medicare", 50.0, ""])
    ws.append(["Carter", "Private Pay", 45.0, "3 paid days, 4th day billed to Parker Grant"])
    buf = io.BytesIO()
    wb.save(buf)

    result = parse_rate_master(buf.getvalue())
    assert len(result.rows) == 2
    baker = next(r for r in result.rows if r.participant_name == "Baker")
    assert baker.payer_source == "Medicare"
    assert baker.grant_rule_type == "none"

    carter = next(r for r in result.rows if r.participant_name == "Carter")
    assert carter.grant_rule_type == "rotating_grant"
    assert carter.grant_cycle_length == 3
    assert "Grant" in (carter.grant_payer or "")


def test_parse_rate_master_prefers_new_sheet_over_old():
    """Real workbooks carry a legacy 'Old Rates' sheet alongside a
    current one; the wrong sheet must never be silently used."""
    wb = openpyxl.Workbook()
    old = wb.active
    old.title = "Old Rates"
    old.append(["Pat ID Number", "Name", "Payor", "Full Rate"])
    old.append(["PAM-1", "Stale, Person", "PP", 50.0])

    new = wb.create_sheet("NEW Rates as of 1.1.2025")
    new.append(["PAT ID Number", "Name", "Payor", "Client Rate", "Discount"])
    new.append([None, "Current, Person", "Medicaid", 94.66, None])

    buf = io.BytesIO()
    wb.save(buf)

    result = parse_rate_master(buf.getvalue())
    assert result.sheet_used == "NEW Rates as of 1.1.2025"
    names = [r.participant_name for r in result.rows]
    assert names == ["Current, Person"]


def test_parse_rate_master_discount_column_recognized_and_normalized():
    """Real header is literally "Discount" (not "notes"/"exception"), and
    the Payor column ("Medacaid", a typo) must be normalized so it can
    be compared against a PCC batch's payer categorization later."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Participant Name", "Payor", "Client Rate", "Discount"])
    ws.append(["Ferrer, Lavinia", "Medacaid", 94.66, None])
    ws.append(["Berkin, Harvey", "Private Pay", 95, "4th day is Parker Grant"])
    buf = io.BytesIO()
    wb.save(buf)

    result = parse_rate_master(buf.getvalue())
    ferrer = next(r for r in result.rows if r.participant_name == "Ferrer, Lavinia")
    assert ferrer.payer_source == "Medicaid"

    berkin = next(r for r in result.rows if r.participant_name == "Berkin, Harvey")
    assert berkin.grant_rule_type == "rotating_grant"
    assert berkin.grant_cycle_length == 3
    assert berkin.grant_cycle_secondary_days == 1
    assert berkin.grant_payer == "Parker Grant"


def test_parse_rate_master_rotation_with_secondary_payer_count():
    """'Three days private, then two days Title III' -- a rotation to a
    named payer OTHER than a grant, with more than one secondary day."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Participant Name", "Payor", "Client Rate", "Discount"])
    ws.append(["Cavalieri, Giovanni", "Private Pay Discount $53.00 Daily", 42, "Three days private, then two days Title III"])
    buf = io.BytesIO()
    wb.save(buf)

    result = parse_rate_master(buf.getvalue())
    row = result.rows[0]
    assert row.payer_source == "Private Pay"
    assert row.grant_rule_type == "rotating_grant"
    assert row.grant_cycle_length == 3
    assert row.grant_cycle_secondary_days == 2
    assert row.grant_payer == "Title III"


def test_parse_rate_master_no_daycount_leaves_rule_none():
    """"Full PG while appealing Medicaid denial" has no day count -- it's
    a status note, not a parseable rotation, and must not be guessed at."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Participant Name", "Payor", "Client Rate", "Discount"])
    ws.append(["Hall, Tresea", "Private Pay", 95, "Full PG while appealing Medicaid denial"])
    buf = io.BytesIO()
    wb.save(buf)

    result = parse_rate_master(buf.getvalue())
    row = result.rows[0]
    assert row.grant_rule_type == "none"
    assert "Full PG while appealing Medicaid denial" in (row.notes or "")


def test_parse_pcc_batch_text_heuristic():
    import fpdf
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in [
        "Adams, Alice   Private Pay   $45.00",
        "Baker, Bob     Medicare      $50.00",
    ]:
        pdf.cell(0, 8, text=line, new_x="LMARGIN", new_y="NEXT")
    raw = bytes(pdf.output())

    result = parse_pcc_batch(raw)
    assert result.extraction_method == "pdf-text-heuristic"
    assert len(result.rows) == 2
    names = {r.raw_name for r in result.rows}
    assert names == {"Adams, Alice", "Baker, Bob"}
    amounts = {r.raw_name: r.amount for r in result.rows}
    assert amounts["Adams, Alice"] == 45.0
    assert amounts["Baker, Bob"] == 50.0
    payers = {r.raw_name: r.payer_source for r in result.rows}
    assert payers["Baker, Bob"] == "Medicare"
